#!/usr/bin/env python3
"""
read_tracker.py — read the "RFC Tracker.xlsx" workbook into a flat JSON list.

The tracker is the authoritative source for the metadata the PDF form does NOT
contain: System Name, Before, After, Executor, Execution Date. It also carries
the "Link Document" column (a filename) which is the join key to the parsed
PDF records.

Each output row keeps the raw "Link Document" string; the matching to actual
PDF filenames (prefix/substring tolerant) happens in the Node ingest step,
which has both sides.

Output: RFC/imported/_tracker.json (list of rows).
"""
import json, os, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "RFC" / "RFC Tracker.xlsx"
OUT  = ROOT / "RFC" / "imported" / "_tracker.json"

# Header labels we care about (matched by substring, case-insensitive).
WANT = {
    "system":      "system name",
    "title":       "rfc title",
    "description": "change desc",
    "before":      "before",
    "after":       "after",
    "execDate":    "execution date",
    "executor":    "executor name",
    "product":     "product name",
    "docLink":     "link af1 sign form",
    "filename":    "link document",
}


def main():
    if not XLSX.exists():
        print(f"tracker not found at {XLSX}", file=sys.stderr)
        sys.exit(1)
    OUT.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        grid = list(ws.iter_rows(values_only=True))
        # locate the header row in this sheet (contains "System Name")
        hdr_idx = None
        for i, r in enumerate(grid):
            cells = [str(c).lower() if c is not None else "" for c in r]
            if any("system name" in c for c in cells):
                hdr_idx = i
                break
        if hdr_idx is None:
            continue  # sheet without the RFC table
        headers = [str(c).lower().strip() if c is not None else "" for c in grid[hdr_idx]]
        col_for = {}
        for key, label in WANT.items():
            for ci, h in enumerate(headers):
                if label in h:
                    col_for[key] = ci
                    break
        for r in grid[hdr_idx + 1:]:
            if not any(c is not None and str(c).strip() for c in r):
                continue  # blank row
            row = {"sheet": sheet}
            for key, ci in col_for.items():
                val = r[ci] if ci < len(r) else None
                if isinstance(val, float) and val.is_integer():
                    val = int(val)
                row[key] = val
            # only keep rows that look like an RFC (have a system or title)
            if row.get("system") or row.get("title") or row.get("filename"):
                rows.append(row)

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2, default=str)
    print(f"wrote {len(rows)} tracker rows → {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
